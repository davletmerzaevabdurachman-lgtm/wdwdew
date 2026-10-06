from pathlib import Path
import csv, json, io

def extract_text(filename: str, data: bytes) -> str:
    ext=Path(filename).suffix.lower()
    if ext in {".txt",".md",".json",".csv",".py",".js",".ts",".tsx",".jsx",".css",".html",".sql",".sh"}:
        return data.decode("utf-8", errors="replace")
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    if ext == ".docx":
        from docx import Document
        return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    if ext in {".xlsx",".xls"}:
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return "\n".join(" | ".join("" if v is None else str(v) for v in row) for ws in wb.worksheets for row in ws.iter_rows(values_only=True))
    raise ValueError(f"Unsupported file type: {ext}")
